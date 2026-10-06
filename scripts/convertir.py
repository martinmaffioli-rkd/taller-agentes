"""Convierte los documentos de conocimiento/originales/ (Word, Excel, PDF, CSV) a texto Markdown
en conocimiento/, que es lo que lee el agente. Uso: python scripts/convertir.py [carpeta]"""
import csv
import sys
from comun import carpeta

LIMITE_KB = 150


def tabla_md(filas):
    filas = [["" if v is None else str(v).replace("|", "/").replace("\n", " ") for v in f] for f in filas]
    filas = [f for f in filas if any(c.strip() for c in f)]
    if not filas:
        return ""
    ancho = max(len(f) for f in filas)
    filas = [f + [""] * (ancho - len(f)) for f in filas]
    lineas = ["| " + " | ".join(filas[0]) + " |", "|" + "---|" * ancho]
    lineas += ["| " + " | ".join(f) + " |" for f in filas[1:]]
    return "\n".join(lineas)


def desde_excel(ruta):
    from openpyxl import load_workbook
    libro = load_workbook(ruta, data_only=True)
    partes = []
    for hoja in libro.worksheets:
        partes.append(f"## Hoja: {hoja.title}\n\n" + tabla_md(hoja.iter_rows(values_only=True)))
    return "\n\n".join(partes)


def desde_word(ruta):
    from docx import Document
    doc = Document(ruta)
    partes = []
    for p in doc.paragraphs:
        texto = p.text.strip()
        if not texto:
            continue
        estilo = (p.style.name or "").lower()
        if estilo.startswith("heading") or estilo.startswith("título") or estilo == "title":
            nivel = "".join(ch for ch in estilo if ch.isdigit()) or "1"
            partes.append("#" * (int(nivel) + 1) + " " + texto)
        elif "list" in estilo:
            partes.append("- " + texto)
        else:
            partes.append(texto)
    for t in doc.tables:
        partes.append(tabla_md([[celda.text for celda in fila.cells] for fila in t.rows]))
    return "\n\n".join(partes)


def desde_pdf(ruta):
    from pypdf import PdfReader
    paginas = [p.extract_text() or "" for p in PdfReader(ruta).pages]
    return "\n\n".join(f"## Página {i}\n\n{t.strip()}" for i, t in enumerate(paginas, 1) if t.strip())


def desde_csv(ruta):
    with open(ruta, newline="", encoding="utf-8") as f:
        return tabla_md(csv.reader(f))


CONVERSORES = {".xlsx": desde_excel, ".docx": desde_word, ".pdf": desde_pdf, ".csv": desde_csv,
               ".md": lambda r: r.read_text(encoding="utf-8"), ".txt": lambda r: r.read_text(encoding="utf-8")}

base = carpeta(sys.argv[1] if len(sys.argv) > 1 else None)
origen, destino = base / "conocimiento" / "originales", base / "conocimiento"
archivos = sorted(p for p in origen.glob("*") if p.is_file()) if origen.exists() else []
if not archivos:
    sys.exit("No hay documentos en conocimiento/originales/.")

for ruta in archivos:
    conv = CONVERSORES.get(ruta.suffix.lower())
    if not conv:
        print(f"⚠️  {ruta.name}: formato no soportado (usá Word, Excel, PDF o CSV)")
        continue
    texto = f"# {ruta.stem}\n\n_Documento convertido desde {ruta.name}._\n\n" + conv(ruta)
    salida = destino / f"{ruta.stem}.md"
    salida.write_text(texto, encoding="utf-8")
    kb = len(texto.encode()) / 1024
    aviso = f"  ⚠️ es grande ({kb:.0f} KB), conviene resumirlo" if kb > LIMITE_KB else ""
    print(f"✅ {ruta.name} → conocimiento/{salida.name} ({kb:.1f} KB){aviso}")
